"""Prepara arquivos locais a partir de evidencias existentes; sem rede/DB/publicacao."""
from __future__ import annotations
import argparse
from collections import Counter
from decimal import Decimal, InvalidOperation
import hashlib
import math
import json
from pathlib import Path
import re
from conteudo_editorial import validate_bundle

ROOT = Path(__file__).resolve().parents[1]
LOCAL = ROOT / "content/local"
REPORT = ROOT / "content/manifests/pilot-prepared.json"
WGER = ["docs/planejamento/evidencias/wger-25.json", "docs/conteudo/evidencias/wger-offset-300.json", "docs/conteudo/evidencias/wger-offset-600.json"]
CANDIDATES = "content/manifests/pilot-candidates.json"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read(name):
    return json.loads((ROOT / name).read_text(encoding="utf-8-sig"))


def numeric_cell(raw):
    result = {"raw": raw, "state": "missing", "value": None, "upperBound": None}
    if raw is None or str(raw).strip() in {"", "-"}:
        return result
    value = str(raw).strip()
    if value.lower() in {"tr", "trace", "traces"}:
        result["state"] = "trace"
        return result
    if value.upper() == "N":
        result["state"] = "present_quantity_unknown"
        return result
    bound = value.startswith("<")
    token = (value[1:] if bound else value).strip().replace(",", ".")
    try:
        number = Decimal(token)
        if not number.is_finite() or number < 0:
            raise InvalidOperation
        number = float(number)
        if not math.isfinite(number):
            raise InvalidOperation
        result["state"] = "less_than" if bound else "numeric"
        result["upperBound" if bound else "value"] = number
    except InvalidOperation:
        result["state"] = "unparsed"
    return result


def pointer(doc, path):
    for part in path.split("/")[1:]:
        part = part.replace("~1", "/").replace("~0", "~")
        doc = doc[int(part)] if isinstance(doc, list) else doc[part]
    return doc


def evidence(path, point):
    return {"path": path, "pointer": point, "snapshotSha256": sha(ROOT / path)}


def asset(ident, kind, source, ref, title=None, locale=None, version=None, authors=None, rights=None, audience=None, content=None, dose=None, reasons=None):
    rights = dict(rights or {"label": None, "licenseUrl": None, "attribution": None, "scope": "movement_metadata"})
    rights["operationGrants"] = {op: "unknown" for op in ("catalog_metadata", "instruction_text", "media_display", "guided_practice", "export_dataset")}
    return {"id": ident, "kind": kind, "sourceId": source, "sourceVersion": version, "locale": locale, "title": title,
            "authors": authors or [], "evidence": [ref],
            "rights": rights or {"label": None, "licenseUrl": None, "attribution": None, "scope": "movement_metadata"},
            "audience": audience or {"population": None, "environment": None, "supportRequired": None, "supportSource": None},
            "dose": dose or {"measurements": [], "sourceFacts": None, "contextRef": None}, "content": content or {},
            "review": {"state": "blocked" if reasons else "pending", "reviewer": None, "rightsVerified": False, "fidelityChecked": False,
                       "safetyUseReviewed": False, "instructionSufficient": False, "reasons": reasons or ["Revisao por ativo/uso pendente; nao seed do aplicativo"]},
            "allowedOperations": []}


def wger_assets():
    documents = [(name, read(name)) for name in WGER]
    licenses, assets, sources, ids = {}, [], [], set()
    languages_doc = read("docs/planejamento/evidencias/wger-languages.json")
    languages = {x["id"]: x["short_name"] for x in languages_doc.get("data", languages_doc)["results"]}
    for name, doc in documents:
        payload = doc.get("data", doc)
        for row in payload["results"]:
            licenses[row["license"]["id"]] = row["license"]
    for name, doc in documents:
        sources.append({"path": name, "sha256": sha(ROOT / name)})
        for index, row in enumerate(doc.get("data", doc)["results"]):
            if row["id"] in ids:
                raise ValueError("Exercicio duplicado em snapshots: " + str(row["id"]))
            ids.add(row["id"])
            base = ("/data" if "data" in doc else "") + f"/results/{index}"
            movement_id = f"wger:movement:{row['id']}"
            translations = row.get("translations", [])
            chosen = next((t for t in translations if t["language"] == 7), next((t for t in translations if t["language"] == 2), translations[0] if translations else {}))
            common_reasons = ["ID1022: traducoes semanticamente divergentes na auditoria anterior"] if row["id"] == 1022 else []
            lic = row["license"]
            rights = {"label": lic["short_name"], "licenseUrl": lic["url"], "scope": "movement_metadata",
                      "attribution": "wger / " + ", ".join(row.get("total_authors_history", [])), "upstreamLicenseAuthor": row.get("license_author")}
            assets.append(asset(movement_id, "movement", "EX01", evidence(name, base), chosen.get("name"), languages.get(chosen.get("language")), None, row.get("total_authors_history"), rights,
                                content={"sourceRecordId": row["id"], "sourceUuid": row["uuid"], "sourceLastUpdateGlobal": row.get("last_update_global"), "category": row["category"], "equipment": row["equipment"],
                                         "primaryMuscles": [{"id": m["id"], "name": m["name"], "name_en": m.get("name_en")} for m in row["muscles"]],
                                         "selectedNameTranslationId": chosen.get("id"), "instructionRef": f"wger:translations:{chosen['id']}" if chosen.get("id") else None, "instructionsAvailable": len(translations), "imageReferencesAvailable": len(row.get("images", [])), "videoReferencesAvailable": len(row.get("videos", []))}, reasons=common_reasons))
            for key, kind in [("translations", "instruction"), ("images", "media"), ("videos", "media")]:
                for j, item in enumerate(row.get(key, [])):
                    childlic = licenses.get(item.get("license"), {})
                    childrights = {"label": childlic.get("short_name"), "licenseUrl": childlic.get("url"), "scope": "instruction_text" if kind == "instruction" else "media_asset",
                                   "attribution": (item.get("license_author") or ", ".join(item.get("author_history", []))) or None,
                                   "licenseId": item.get("license"), "objectUrl": item.get("license_object_url"), "derivativeSourceUrl": item.get("license_derivative_source_url")}
                    content = {"movementRef": movement_id, "sourceObjectUuid": item.get("uuid"), "sourceLanguageId": item.get("language"), "sourceCreated": item.get("created")}
                    reasons = list(common_reasons)
                    if kind == "instruction":
                        content.update(originalDescription=item.get("description"), originalDescriptionSource=item.get("description_source"), originalHtmlPreserved=True, translationCreatedByThisTool=False)
                    else:
                        content.update(mediaType="image" if key == "images" else "video", remoteUrl=item.get("image", item.get("video")), copied=False, aiGenerated=item.get("is_ai_generated"), metadataOnly=True)
                        if item.get("is_ai_generated") is True:
                            reasons.append("Midia sinalizada IA; nao usar como demonstracao humana validada")
                    if not childrights["attribution"]:
                        reasons.append("Credito por objeto ausente no snapshot; nao herdar autor do movimento")
                    if not childrights["licenseUrl"]:
                        reasons.append("Licenca do objeto nao resolvida no snapshot")
                    assets.append(asset(f"wger:{key}:{item['id']}", kind, "EX01", evidence(name, base + f"/{key}/{j}"), item.get("name"), languages.get(item.get("language")), None, item.get("author_history"), childrights, content=content, reasons=reasons))
    sources.append({"path": "docs/planejamento/evidencias/wger-languages.json", "sha256": sha(ROOT / "docs/planejamento/evidencias/wger-languages.json")})
    return assets, sources


def selected_assets():
    manifest = read(CANDIDATES)
    assets, grouped, inputs = [], {}, {}
    for component in manifest["selectedDocumentaryComponents"]:
        path, point = component["evidence"], component["pointer"]
        doc = read(path)
        original = pointer(doc, point)
        inputs[path] = sha(ROOT / path)
        mid = "research:movement:" + component["componentId"]
        context = component.get("contextPointer")
        population = component.get("population")
        audience = {"population": population, "environment": doc.get("identity", {}).get("environment", doc.get("setting")), "supportRequired": None, "supportSource": component.get("delivery"), "selectedResearchContexts": component["niches"]}
        facts = original.get("dose")
        measures = []
        for field, unit in [("holdSeconds", "s"), ("repetitionsPerLeg", "repetitions"), ("stepHeightCm", "cm")]:
            if isinstance(facts, dict) and field in facts:
                measures.append({"quantity": facts[field], "unit": unit, "sourceField": point + "/dose/" + field})
        dose = {"measurements": measures, "sourceFacts": facts, "contextRef": {"path": path, "pointer": context} if context else None,
                "sourceDoseReference": original.get("doseRef", original.get("doseSource")), "noProgrammeDoseTransferredToMovement": True}
        rights = {"label": component.get("knownRights"), "licenseUrl": None, "attribution": None, "scope": "movement_metadata"}
        assets.append(asset(mid, "movement", component["sourceId"], evidence(path, point), component["name"], "pt", doc.get("identity", {}).get("edition"), component.get("authors"), rights, audience, {"originalComponentId": component["componentId"], "instructionRef": "research:instruction:" + component["componentId"], "mediaCopied": False}, dose))
        text_facts = {k:original[k] for k in ("preparation", "movement", "ending", "instructionPtDraft", "executionAndReturn", "methodFactsInPortuguese", "returnEvidence", "exactVariant", "support", "limits") if k in original}
        text_rights = dict(rights, scope="instruction_text")
        assets.append(asset("research:instruction:" + component["componentId"], "instruction", component["sourceId"], evidence(path, point), component["name"], "pt", doc.get("identity", {}).get("edition"), component.get("authors"), text_rights, audience,
                            {"movementRef": mid, "documentaryFacts": text_facts, "sourceFidelityReviewRequired": True, "textCreatedByThisTool": False}))
        key = (path, original.get("group", "package"))
        grouped.setdefault(key, []).append(mid)
    for (path, group), members in grouped.items():
        doc = read(path)
        is_home = doc.get("id") == "ES2019-DOM-01"
        kind = "program" if is_home else "collection"
        source = next(x["sourceId"] for x in manifest["selectedDocumentaryComponents"] if x["evidence"] == path)
        payload = {"itemRefs": members, "inheritsSourceEfficacy": False, "noExecutionSequenceCreated": not is_home, "orderIsDocumentaryIndexOnly": not is_home, "newProgrammeCreated": False}
        if is_home:
            payload.update(sourceOrderAndDoseNotChanged=True, sourceProtocolRef={"path": path, "pointer": "/dose"}, sourceDoseContext=doc["dose"], sourceProtocolOnly=True)
        elif group in doc.get("sourceGroups", {}):
            source = doc["sourceGroups"][group]["sourceId"]
            payload["sourceContextRef"] = {"path": path, "pointer": "/sourceGroups/" + group}
        context = doc.get("sourceGroups", {}).get(group, doc)
        audience = {"population": context.get("population"), "environment": doc.get("identity", {}).get("environment", doc.get("setting")), "supportRequired": None,
                    "supportSource": context.get("teachingAndSupport", context.get("delivery"))}
        authors = doc.get("identity", {}).get("authors", context.get("authors", []))
        original_rights = doc.get("rights", {}) if is_home else {}
        assets.append(asset("research:" + kind + ":" + Path(path).parent.name + ":" + group, kind, source, evidence(path, ""), title=doc.get("id", group), locale="pt", version=doc.get("identity", {}).get("edition"), authors=authors, audience=audience, content=payload,
                            rights={"label": original_rights.get("license"), "licenseUrl": original_rights.get("licenseUrlReportedInLocalArticle"), "attribution": original_rights.get("attributionDraft"), "scope": "source_protocol" if is_home else "modular_collection"}))
    return assets, [{"path": p, "sha256": h} for p,h in inputs.items()]


def prepare_editorial():
    catalog = {x["id"]:x for x in read("docs/conteudo/fontes.json")["sources"]}
    wa, wi = wger_assets()
    sa, si = selected_assets()
    assets = wa + sa
    sources = [{"id": s, "name": catalog[s]["name"], "publisher": catalog[s].get("publisher"), "evidence": ["docs/conteudo/fontes.json"], "urls": catalog[s].get("urls", [])} for s in sorted({a["sourceId"] for a in assets})]
    bundle = {"schemaVersion": "1.0.0", "bundleId": "TF-LOCAL-EDITORIAL-20261021-v1", "stage": "candidate", "sources": sources, "assets": assets, "runtimeSeed": False, "inputs": wi + si + [{"path": p, "sha256": sha(ROOT / p)} for p in ["docs/conteudo/fontes.json", CANDIDATES]]}
    errors = validate_bundle(bundle)
    if errors:
        raise ValueError("; ".join(errors[:10]))
    target = LOCAL / "editorial-v1.json"
    target.write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    matrix = {"schemaVersion": "1.0.0", "purpose": "Matriz documental sem texto instrucional nem midia", "stage": "candidate", "sourceBundleSha256": sha(target), "sourceBodiesIncluded": False,
              "rows": [{"assetId": a["id"], "kind": a["kind"], "sourceId": a["sourceId"], "sourceVersion": a["sourceVersion"], "evidence": a["evidence"],
                        "licenseLabel": a["rights"]["label"], "licenseUrl": a["rights"]["licenseUrl"], "attributionPresent": bool(a["rights"]["attribution"]),
                        "operationGrants": a["rights"]["operationGrants"], "allowedOperations": a["allowedOperations"], "reviewState": a["review"]["state"], "reasons": a["review"]["reasons"]} for a in assets]}
    matrix_path = ROOT / "content/manifests/pilot-rights-matrix.json"
    matrix_path.write_text(json.dumps(matrix,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    return {"path": target.relative_to(ROOT).as_posix(), "sha256": sha(target), "rightsMatrix": {"path": matrix_path.relative_to(ROOT).as_posix(), "sha256": sha(matrix_path)}, "counts": dict(Counter(a["kind"] for a in assets)), "blockedAssets": sum(a["review"]["state"] == "blocked" for a in assets),
            "portugueseWgerInstructions": sum(a["sourceId"] == "EX01" and a["kind"] == "instruction" and a["locale"] == "pt" for a in assets), "mediaDownloaded": 0, "runtimeApproved": 0, "inputs": wi + si + [{"path": p, "sha256": sha(ROOT / p)} for p in ["docs/conteudo/fontes.json", CANDIDATES]]}


def prepare_food(name, receipt_name, source_id, sheet_names=None):
    import openpyxl
    receipt = read(receipt_name)
    original = ROOT / name
    if receipt.get("status") != 200 or receipt.get("sha256") != sha(original):
        raise ValueError("Recibo/hash da planilha invalido: " + name)
    workbook = openpyxl.load_workbook(original, read_only=True, data_only=True)
    target = LOCAL / (source_id.lower() + "-foods.jsonl")
    states, unique, records, rejected = Counter(), set(), 0, 0
    food_names, identity_conflicts = set(), []
    sheets = sheet_names or workbook.sheetnames[2:]
    source_columns = {}
    with target.open("w", encoding="utf-8") as f:
        for name_sheet in sheets:
            sheet = workbook[name_sheet]
            rows = sheet.iter_rows(values_only=True)
            header = next(rows)
            ciqual = source_id == "AL-CIQUAL-2025"
            if ciqual:
                start, offset = 2, 9
                code_i, name_i = header.index("alim_code"), header.index("alim_nom_fr")
            else:
                next(rows); next(rows)
                start, offset, code_i, name_i = 4, 7, 0, 1
            body_rows = list(enumerate(rows, start))
            by_code = {}
            for row_index, source_row in body_rows:
                if source_row[code_i] is not None and source_row[name_i] is not None:
                    by_code.setdefault(str(source_row[code_i]).strip(), []).append((row_index, source_row[name_i]))
            conflicts = {code: values for code, values in by_code.items() if len(values) > 1}
            identity_conflicts += [{"sheet": name_sheet, "sourceFoodId": code, "rowsAndNames": values} for code, values in conflicts.items()]
            columns = [(i,str(h)) for i,h in enumerate(header) if h is not None and i >= offset]
            source_columns[name_sheet] = [{"column": i + 1, "header": h} for i,h in columns]
            for index,row in body_rows:
                code = row[code_i]
                if code is None or row[name_i] is None:
                    rejected += 1
                    continue
                code = str(code).strip()
                key = (name_sheet, code)
                unique.add(key)
                food_names.add((code, row[name_i]))
                values = []
                for col, text in columns:
                    v = numeric_cell(row[col])
                    states[v["state"]] += 1
                    values.append({"column": col + 1, **v})
                if ciqual:
                    basis = {"quantity": 100, "unit": "g", "kind": "food", "source": "Cabecalhos originais da tabela"}
                elif "100gFA" in name_sheet or "100FA" in name_sheet:
                    basis = {"quantity": 100, "unit": "g", "kind": "fatty_acids", "source": "Nome da tabela; NAO por100g de alimento"}
                elif "100gFood" in name_sheet:
                    basis = {"quantity": 100, "unit": "g", "kind": "food", "source": "Nome da tabela"}
                elif name_sheet == "1.2 Factors":
                    basis = {"quantity": None, "unit": None, "kind": "factor", "source": "Fatores nao sao nutrientes por100g"}
                else:
                    basis = {"quantity": 100, "unit": None, "kind": "food_or_alcohol_exception", "source": "1.1 Notes!A4:100g exceto bebidas alcoolicas100ml; resolver grupo antes do calculo"}
                record = {"sourceId": source_id, "sourceVersion": "2025" if ciqual else "2021", "sourceFoodId": code, "sourceRowKey": f"{source_id}:{name_sheet}:{index}", "sourceCodeConflict": code in conflicts,
                          "nameOriginal": row[name_i], "nameLocale": "fr" if ciqual else "en", "namePt": None,
                          "sheet": name_sheet, "row": index, "quantityBasis": basis, "values": values,
                          "sourceContext": list(row[:offset]), "calculationsApproved": False, "runtimeImportApproved": False}
                f.write(json.dumps(record,ensure_ascii=False,separators=(",", ":")) + "\n")
                records += 1
    workbook.close()
    return {"path": target.relative_to(ROOT).as_posix(), "sha256": sha(target), "sourceId": source_id, "sourceFile": name, "sourceSha256": sha(original),
            "receipt": receipt_name, "receiptSha256": sha(ROOT / receipt_name), "sourceUrl": receipt["url"], "tables": len(sheets), "uniqueFoods": len(food_names), "uniqueSourceCodes": len({code for _,code in unique}), "tableRecords": records, "identityConflicts": identity_conflicts,
            "blankOrNamelessRows": rejected, "valueStates": dict(states), "columnsByTable": source_columns,
            "credits": "Anses, Table Ciqual2025, doi:10.57745/RDMHWY, Licence Ouverte Etalab2.0" if ciqual else "Public Health England, CoFID2021, Open Government Licence v3.0, excecoes aplicaveis",
            "originalLabelsPreserved": True, "portugueseTranslationReviewed": False, "runtimeApproved": 0}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true")
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    if args.write:
        LOCAL.mkdir(parents=True,exist_ok=True)
        report = {"schemaVersion": 1, "date": "2026-10-21", "stage": "local_candidate", "noNetworkCalls": True, "noDatabaseWrites": True,
                  "editorial": prepare_editorial(),
                  "food": [prepare_food("docs/conteudo/evidencias/ciqual-2025.xlsx", "docs/conteudo/evidencias/ciqual-2025-windows.evidencia.json", "AL-CIQUAL-2025", ["composition nutritionnelle"]),
                           prepare_food("docs/conteudo/evidencias/cofid-2021.xlsx", "docs/conteudo/evidencias/cofid-2021.evidencia.json", "AL-COFID-2021")],
                  "programmeOrClinicalApproval": False, "operationalApproval": False,
                  "toolInputs": [{"path": p, "sha256": sha(ROOT / p)} for p in ["tools/preparar_conteudo_local.py", "tools/conteudo_editorial.py", "content/schemas/editorial-v1.schema.json"]]}
        REPORT.write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    if args.check:
        report = read(REPORT.relative_to(ROOT))
        for product in [report["editorial"], *report["food"]]:
            path = ROOT / product["path"]
            if sha(path) != product["sha256"]:
                raise ValueError("Conteudo local alterado: " + product["path"])
        bundle = read(report["editorial"]["path"])
        errors = validate_bundle(bundle)
        if errors:
            raise ValueError("; ".join(errors[:10]))
        matrix_info = report["editorial"]["rightsMatrix"]
        if sha(ROOT / matrix_info["path"]) != matrix_info["sha256"]:
            raise ValueError("Matriz de direitos alterada")
        matrix = read(matrix_info["path"])
        if matrix["sourceBundleSha256"] != report["editorial"]["sha256"] or {x["assetId"] for x in matrix["rows"]} != {a["id"] for a in bundle["assets"]}:
            raise ValueError("Matriz de direitos fora de sincronia")
        for entry in report["editorial"]["inputs"]:
            if sha(ROOT / entry["path"]) != entry["sha256"]:
                raise ValueError("Entrada editorial alterada: " + entry["path"])
        for tool in report.get("toolInputs", []):
            if sha(ROOT / tool["path"]) != tool["sha256"]:
                raise ValueError("Ferramenta/schema mudou; reconstruir o lote")
        for food in report["food"]:
            if sha(ROOT / food["receipt"]) != food["receiptSha256"]:
                raise ValueError("Recibo alimentar alterado")
            if sha(ROOT / food["sourceFile"]) != food["sourceSha256"]:
                raise ValueError("Entrada alimentar alterada")
        print(json.dumps({"result": "PASS", "editorial": report["editorial"]["counts"], "food": [{k:f[k] for k in ["sourceId", "tables", "uniqueFoods", "tableRecords"]} for f in report["food"]], "runtimeApproved": 0},ensure_ascii=True))
    if not args.write and not args.check:
        parser.error("Use --write ou --check; nao ha acesso a rede ou banco")


if __name__ == "__main__":
    main()
